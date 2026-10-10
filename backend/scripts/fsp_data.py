"""
Реальные результаты соревнований ФСП: шаблон таблицы, проверка, обновление профилей.

    python -m scripts.fsp_data template   # создать data/fsp_results_template.xlsx (с инструкцией и примерами)
    python -m scripts.fsp_data check      # проверить data/fsp_results.xlsx: сколько участников, какие строки с ошибками
    python -m scripts.fsp_data sync       # обновить достижения у кандидатов, которые уже привязали свой ФСП ID

Как пользоваться:
  1. template -> открыть файл в Excel, заполнить лист «Данные» (лист «Инструкция» подскажет, что куда).
  2. Сохранить как data/fsp_results.xlsx (можно и .csv — тогда указать FSP_DATA_PATH=data/fsp_results.csv).
  3. check -> исправить строки, на которые он ругается.
  4. Бэкенд подхватывает файл сам, без перезапуска. Кандидат, который привяжет ID из таблицы,
     получит достижения с пометкой «verified». Уже привязанным — обновить через sync.

В Docker: docker compose exec backend python -m scripts.fsp_data check
"""
import sys
from pathlib import Path

from app.core.config import settings
from app.services.fsp_data import COLUMNS, load_dataset
from app.services.fsp_service import BACKEND_DIR, FileFspRegistry

TEMPLATE_PATH = BACKEND_DIR / "data" / "fsp_results_template.xlsx"

HINTS = {
    "fsp_id": ("ID участника ФСП", "Обязательно. Одинаковый у всех строк одного человека", "100245"),
    "event": ("Соревнование", "Пусто = участник есть, но достижений нет", "Чемпионат России по спортивному программированию"),
    "discipline": ("Дисциплина", "Необязательно", "Продуктовое программирование"),
    "level": ("Уровень", "Выбрать из списка: международный, всероссийский, региональный, муниципальный", "всероссийский"),
    "year": ("Год", "Число, например 2024", "2024"),
    "place": ("Место", "1, 2, 3 или пусто, если без призового места", "2"),
    "result": ("Результат", "Необязательно: «Финалист», «Участник»… Пусто — возьмём из места", ""),
    "team": ("Команда", "Необязательно", "Команда «Байт»"),
    "sport_rank": ("Разряд", "Необязательно: 3/2/1 разряд, КМС, МС, МСМК. Достаточно указать в одной строке участника",
                   "КМС"),
    "source_url": ("Ссылка на протокол", "Где опубликованы результаты. Очень желательно: это и есть подтверждение", "https://fsp-russia.com/..."),
    "comment": ("Заметки", "Для себя, в систему не попадает", ""),
}


def make_template(path: Path = TEMPLATE_PATH) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(COLUMNS)
    head = PatternFill("solid", fgColor="6D28D9")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    examples = [
        ["100245", "Чемпионат России по спортивному программированию", "Продуктовое программирование",
         "всероссийский", 2024, 2, "", "Команда «Байт»", "КМС", "https://example.org/protocol-2024", "пример — удалить"],
        ["100245", "Региональный этап Чемпионата России", "Продуктовое программирование",
         "региональный", 2024, 1, "", "Команда «Байт»", "", "https://example.org/region-2024", "пример — удалить"],
        ["100377", "Кубок России по спортивному программированию", "Алгоритмическое программирование",
         "всероссийский", 2025, "", "Финалист", "", "1 разряд", "https://example.org/cup-2025", "пример — удалить"],
        ["100500", "", "", "", "", "", "", "", "", "", "пример: участник без достижений — удалить"],
    ]
    for row in examples:
        ws.append(row)
    widths = [12, 48, 34, 18, 8, 8, 14, 20, 12, 40, 30]
    for col, w in zip("ABCDEFGHIJK", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"международный,всероссийский,региональный,муниципальный"',
                        allow_blank=True, showErrorMessage=True, errorTitle="Уровень",
                        error="Выберите уровень из списка")
    ws.add_data_validation(dv)
    dv.add("D2:D2000")
    dv_rank = DataValidation(type="list", formula1='"3 разряд,2 разряд,1 разряд,КМС,МС,МСМК,ЗМС"', allow_blank=True)
    ws.add_data_validation(dv_rank)
    dv_rank.add("I2:I2000")

    info = wb.create_sheet("Инструкция")
    info.append(["Колонка", "Что это", "Как заполнять", "Пример"])
    for cell in info[1]:
        cell.font = Font(bold=True)
    for col in COLUMNS:
        title, how, example = HINTS[col]
        info.append([col, title, how, example])
    info.append([])
    for line in [
        "Одна строка = одно достижение. У участника с тремя соревнованиями — три строки с одним fsp_id.",
        "Названия колонок на листе «Данные» не менять: по ним бэкенд понимает, что где.",
        "ФИО, телефоны и почту НЕ вносить: для связи с профилем кандидата достаточно ID (152-ФЗ).",
        "Строки с примерами (комментарий «пример — удалить») удалить.",
        "Готовый файл сохранить как backend/data/fsp_results.xlsx и выполнить: python -m scripts.fsp_data check",
    ]:
        info.append([line])
    for col, w in zip("ABCD", [14, 22, 70, 50]):
        info.column_dimensions[col].width = w
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def check(path: Path) -> int:
    if not path.exists():
        print(f"Файла {path} нет. Создайте шаблон: python -m scripts.fsp_data template")
        return 1
    ds = load_dataset(path)
    with_ach = sum(1 for a in ds.participants.values() if a)
    print(f"Файл: {path}")
    print(f"Строк с данными: {ds.rows}")
    print(f"Участников: {len(ds.participants)} (с достижениями: {with_ach}, без: {len(ds.participants) - with_ach})")
    print(f"Достижений: {ds.achievements_count}")
    no_source = sum(1 for a in ds.participants.values() for x in a if not x["source_url"])
    if no_source:
        print(f"Без ссылки на протокол: {no_source} — желательно добавить source_url")
    if ds.errors:
        print(f"\nОшибки ({len(ds.errors)}), эти строки НЕ загружены:")
        for e in ds.errors:
            print("  - " + e)
        return 1
    print("\nОшибок нет ✔")
    return 0


def sync() -> int:
    """Обновляет достижения у кандидатов, чей fsp_id уже есть в базе."""
    from app.core.database import SessionLocal, utcnow
    from app.models import CandidateProfile
    from app.services import fsp_service

    db = SessionLocal()
    updated = 0
    try:
        for c in db.query(CandidateProfile).filter(CandidateProfile.fsp_id.isnot(None)):
            participant = fsp_service.registry.get_participant(c.fsp_id)
            if participant is None:
                continue
            c.fsp_achievements = participant["achievements"]
            c.fsp_score = fsp_service.compute_fsp_score(participant["achievements"])
            c.fsp_rank = participant.get("sport_rank")
            c.fsp_synced_at = utcnow()
            updated += 1
        db.commit()
    finally:
        db.close()
    print(f"Обновлено профилей: {updated}")
    return 0


def main(argv: list[str]) -> int:
    cmd = argv[0] if argv else ""
    if cmd == "template":
        print(f"Шаблон создан: {make_template()}")
        return 0
    if cmd == "check":
        path = Path(argv[1]) if len(argv) > 1 else FileFspRegistry(settings.fsp_data_path).path
        return check(path)
    if cmd == "sync":
        return sync()
    print(__doc__)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
