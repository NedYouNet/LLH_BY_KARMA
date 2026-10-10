"""
Чтение реальных результатов соревнований ФСП из таблицы (xlsx или csv).

Открытого API у ФСП нет, поэтому настоящие данные мы собираем вручную в таблицу:
одна строка = одно достижение одного участника. Шаблон таблицы с пояснениями:

    python -m scripts.fsp_data template      # создаст data/fsp_results_template.xlsx

Колонки (порядок не важен, регистр в названиях не важен):
    fsp_id      — ID участника ФСП (обязательно)
    event       — название соревнования (пусто = участник есть, но достижений нет)
    discipline  — дисциплина, например «Продуктовое программирование»
    level       — международный | всероссийский | региональный | муниципальный
    year        — год проведения (2015–текущий)
    place       — 1, 2, 3 или пусто (нет призового места)
    result      — текст результата: «Финалист», «Участник»… (если пусто — возьмём из place)
    team        — название команды (необязательно)
    sport_rank  — спортивный разряд/звание участника: 3/2/1 разряд, КМС, МС… (необязательно, достаточно в одной строке)
    source_url  — ссылка на протокол/страницу с результатами (желательно — это и есть «проверка»)
    comment     — заметки для себя, в систему не попадают

Персональные данные (ФИО, телефоны, почту) в таблицу НЕ кладём: для связи с профилем достаточно ID.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path


LEVEL_NAME = {"international": "Международный", "all_russian": "Всероссийский", "regional": "Региональный",
              "municipal": "Муниципальный"}

COLUMNS = ["fsp_id", "event", "discipline", "level", "year", "place", "result", "team", "sport_rank", "source_url",
           "comment"]

# Спортивные разряды и звания (по возрастанию). Принимаем и сокращения, и полные названия.
SPORT_RANKS = ["3 юношеский разряд", "2 юношеский разряд", "1 юношеский разряд", "3 разряд", "2 разряд", "1 разряд",
               "КМС", "МС", "МСМК", "ЗМС"]
_RANK_ALIASES = {r.lower(): r for r in SPORT_RANKS} | {
    "кандидат в мастера спорта": "КМС", "мастер спорта": "МС", "мастер спорта россии": "МС",
    "мастер спорта международного класса": "МСМК", "заслуженный мастер спорта": "ЗМС",
    "1": "1 разряд", "2": "2 разряд", "3": "3 разряд", "i": "1 разряд", "ii": "2 разряд", "iii": "3 разряд",
    "1 взрослый разряд": "1 разряд", "2 взрослый разряд": "2 разряд", "3 взрослый разряд": "3 разряд",
}
REQUIRED = ["fsp_id"]

# Принимаем и коды, и русские слова в любом регистре
_LEVEL_ALIASES = {
    "international": "international", "международный": "international", "международные": "international",
    "all_russian": "all_russian", "всероссийский": "all_russian", "всероссийские": "all_russian",
    "российский": "all_russian", "россия": "all_russian",
    "regional": "regional", "региональный": "regional", "региональные": "regional",
    "межрегиональный": "regional",
    "municipal": "municipal", "муниципальный": "municipal", "городской": "municipal",
    "муниципальные": "municipal",
}


@dataclass
class FspDataset:
    participants: dict[str, list[dict]] = field(default_factory=dict)  # fsp_id -> достижения
    ranks: dict[str, str] = field(default_factory=dict)                # fsp_id -> разряд
    errors: list[str] = field(default_factory=list)                   # «Строка 7: ...»
    rows: int = 0

    @property
    def achievements_count(self) -> int:
        return sum(len(a) for a in self.participants.values())


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():  # Excel хранит 2024 как 2024.0
        value = int(value)
    return str(value).strip()


def _read_rows(path: Path) -> list[dict[str, str]]:
    """Строки таблицы как словари {колонка: значение}. Поддерживает .xlsx и .csv (; или ,)."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook  # импорт здесь: библиотека нужна только для xlsx
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Данные"] if "Данные" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [_clean(h).lower() for h in next(it, [])]
        rows = [dict(zip(header, (_clean(v) for v in r))) for r in it]
        wb.close()
        return rows
    if suffix == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251"):  # Excel в русской Windows сохраняет CSV в cp1251
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Не удалось прочитать CSV: сохраните его в кодировке UTF-8")
        delimiter = ";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        reader.fieldnames = [_clean(h).lower() for h in (reader.fieldnames or [])]
        return [{k: _clean(v) for k, v in r.items() if k} for r in reader]
    raise ValueError(f"Неподдерживаемый формат {suffix!r}: нужен .xlsx или .csv")


def parse_level(value: str) -> str | None:
    return _LEVEL_ALIASES.get(value.strip().lower().replace("ё", "е")) if value else None


def parse_rank(value: str) -> str | None | bool:
    """'кмс' -> 'КМС'; пусто -> None; нераспознанное -> False."""
    v = " ".join(value.strip().lower().replace("ё", "е").split())
    if not v:
        return None
    return _RANK_ALIASES.get(v, False)


def parse_place(value: str) -> int | None | str:
    """'1', '1 место', '2-е' -> 1/2; пусто -> None; что-то странное -> 'bad'."""
    v = value.strip().lower()
    if not v:
        return None
    digits = "".join(ch for ch in v.split()[0] if ch.isdigit())
    if digits and int(digits) in (1, 2, 3):
        return int(digits)
    return "bad"


def load_dataset(path: str | Path) -> FspDataset:
    """Читает и проверяет таблицу. Ошибочные строки пропускаются и попадают в errors — с номером строки."""
    path = Path(path)
    ds = FspDataset()
    rows = _read_rows(path)
    if rows and "fsp_id" not in rows[0]:
        ds.errors.append(f"В таблице нет колонки fsp_id. Найдены колонки: {', '.join(rows[0])}")
        return ds
    this_year = date.today().year
    for n, row in enumerate(rows, start=2):  # строка 1 — заголовок, как в Excel
        if not any(row.get(c) for c in COLUMNS if c != "comment"):
            continue  # пустая строка
        ds.rows += 1
        if "пример" in row.get("comment", "").lower():
            ds.errors.append(f"Строка {n}: это строка-пример из шаблона — удалите её")
            continue
        fsp_id = row.get("fsp_id", "")
        problems = []
        if not fsp_id:
            problems.append("не заполнен fsp_id")
        elif len(fsp_id) > 50:
            problems.append("fsp_id длиннее 50 символов")
        rank = parse_rank(row.get("sport_rank", ""))
        if rank is False:
            problems.append(f"разряд «{row.get('sport_rank')}» не распознан (например: 1 разряд, КМС, МС)")
        elif rank and fsp_id:
            ds.ranks.setdefault(fsp_id, rank)
        event = row.get("event", "")
        if not event:  # участник без достижений — допустимо
            if fsp_id and not problems:
                ds.participants.setdefault(fsp_id, [])
            else:
                ds.errors.append(f"Строка {n}: " + "; ".join(problems))
            continue
        level = parse_level(row.get("level", ""))
        if level is None:
            problems.append(f"уровень «{row.get('level', '')}» не распознан "
                            "(нужно: международный, всероссийский, региональный или муниципальный)")
        year_raw = row.get("year", "")
        year = int(year_raw) if year_raw.isdigit() else None
        if year is None or not 2015 <= year <= this_year:
            problems.append(f"год «{year_raw}» должен быть числом от 2015 до {this_year}")
        place = parse_place(row.get("place", ""))
        if place == "bad":
            problems.append(f"место «{row.get('place')}» — укажите 1, 2, 3 или оставьте пустым")
        url = row.get("source_url", "")
        if url and not url.startswith(("http://", "https://")):
            problems.append("source_url должен начинаться с http:// или https://")
        if problems:
            ds.errors.append(f"Строка {n}: " + "; ".join(problems))
            continue
        ds.participants.setdefault(fsp_id, []).append({
            "event": event,
            "discipline": row.get("discipline") or None,
            "level": level,
            "level_name": LEVEL_NAME[level],
            "year": year,
            "place": place,
            "result": row.get("result") or (f"{place} место" if place else "Участник"),
            "team": row.get("team") or None,
            "source": "fsp_results_table",
            "source_url": url or None,
            "verification_status": "verified",
            "verified": True,
        })
    for ach in ds.participants.values():
        ach.sort(key=lambda a: (-a["year"], a["place"] or 9))
    return ds
